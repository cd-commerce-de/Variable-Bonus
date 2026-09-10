"""
Extracts real per-month Good/Better/Best targets from the CURRENT
(separate-file) scorecard format:

  - <BM Scorecard>.xlsx, sheet 'Brands'  -> Launch Manager (LM (F3M) =
    Germany, Expansion (F3M) = Pan-EU) and Brand Manager (per brand,
    per stage: PY1/Y1/Discontinued)
  - <Leadership Scorecard>.xlsx, sheet 'Leadership', rows 13+ -> R&D
    (revenue only; margin left blank, no margin target exists for R&D)

Replaces extract_monthly_targets.py, which assumed both scorecards lived
as sheets inside one combined workbook named 'BM Scorecard 3' /
'Leadership Scorecard 3' -- that combined-workbook format no longer
matches what's actually being provided (two separate files, sheet named
'Brands' / 'Leadership').

Only writes a month's file if that month has real (non-placeholder)
data in the source -- these scorecards are live trackers that show
0/blank for months not yet reached, and a month with no file here
correctly falls back to the app's own quarterly-estimate logic instead
of writing over it with zeros.

Usage:
  python3 scripts/extract_monthly_targets_v2.py \
      --bm 2026_BM_Scorecard.xlsx --leadership 2026_Leadership_Scorecard.xlsx \
      --out-dir public/targets_monthly
"""
import json
import re
import argparse
from datetime import datetime
from openpyxl import load_workbook


def find_month_columns(ws, header_row=3):
    """Scan a header row for GOOD/BETTER/BEST triples and figure out which
    calendar month each belongs to, using the nearest date cell that
    follows (the block's OWN leading date label is not reliable -- the
    weekly section immediately after each triple is)."""
    max_col = ws.max_column
    triples = []
    c = 1
    while c <= max_col - 2:
        v = ws.cell(row=header_row, column=c).value
        if v == 'GOOD' and ws.cell(row=header_row, column=c + 1).value == 'BETTER' and ws.cell(row=header_row, column=c + 2).value == 'BEST':
            triples.append(c)
            c += 3
        else:
            c += 1

    result = {}
    for good_col in triples:
        month_num = None
        for c2 in range(good_col, min(good_col + 15, max_col + 1)):
            for r2 in (1, 2, 3):
                v = ws.cell(row=r2, column=c2).value
                if isinstance(v, datetime):
                    month_num = v.month
                    break
            if month_num:
                break
        if month_num:
            result[month_num] = {'good': good_col, 'better': good_col + 1, 'best': good_col + 2}
    return result


def gbb(ws, row, cols):
    if row is None:
        return {'good': None, 'better': None, 'best': None}
    return {
        'good': ws.cell(row=row, column=cols['good']).value,
        'better': ws.cell(row=row, column=cols['better']).value,
        'best': ws.cell(row=row, column=cols['best']).value,
    }


def has_real_data(triple):
    """A month is 'real' if at least one of good/better/best is a
    non-zero number -- the live tracker shows 0/0/0 or 0/0/None for
    months not yet reached."""
    for v in triple.values():
        if isinstance(v, (int, float)) and v != 0:
            return True
    return False


def extract_bm_scorecard(ws, cols):
    """LM(F3M)/Expansion(F3M) + per-brand PY1/Y1/Discontinued sections
    from the 'Brands' sheet."""
    labels = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v:
            labels[r] = str(v).strip()
        if v == 'Total (Amazon EU)':
            break  # stop before grand-total sections, not needed
    rows_sorted = sorted(labels.keys())

    def find_after(header_row, contains):
        for r in rows_sorted:
            if r <= header_row:
                continue
            if re.match(r'^(LM|Expansion|BM\d)', labels[r]) and r != header_row:
                break  # hit the next section header before finding it
            if contains.lower() in labels[r].lower():
                return r
        return None

    lm_header = next(r for r in rows_sorted if labels[r] == 'LM (F3M)')
    exp_header = next(r for r in rows_sorted if labels[r] == 'Expansion (F3M)')
    lm_rev_row = find_after(lm_header, 'Rolling F3M Revenue')
    lm_margin_row = find_after(lm_header, 'Rolling F3M Profit')
    exp_rev_row = find_after(exp_header, 'Rolling F3M Revenue')
    exp_margin_row = find_after(exp_header, 'Rolling F3M Profit')

    launch_manager = {
        'germany': {'revenue': gbb(ws, lm_rev_row, cols), 'profit_margin': gbb(ws, lm_margin_row, cols)},
        'pan_eu': {'revenue': gbb(ws, exp_rev_row, cols), 'profit_margin': gbb(ws, exp_margin_row, cols)},
    }

    section_re = re.compile(r'^BM\d+ - (.+?) \((PY1|Y1|Discontinued|Total)\)$')
    sections = []
    for r in rows_sorted:
        m = section_re.match(labels[r])
        if m:
            sections.append((r, m.group(1).strip(), m.group(2)))

    brand_manager = {}
    for i, (r, brand, stage) in enumerate(sections):
        if stage not in ('PY1', 'Y1', 'Discontinued'):
            continue
        next_row = sections[i + 1][0] if i + 1 < len(sections) else (rows_sorted[-1] + 1)
        rev_row = margin_row = None
        for rr in rows_sorted:
            if r < rr < next_row:
                lbl = labels[rr].lower()
                if 'revenue' in lbl and 'growth' not in lbl and rev_row is None:
                    rev_row = rr
                if 'profit margin' in lbl and margin_row is None:
                    margin_row = rr
        stage_label = {'PY1': 'PY1', 'Y1': 'Y1 (F4-12)', 'Discontinued': 'Discontinued'}[stage]
        brand_manager.setdefault(brand, {})[stage_label] = {
            'revenue': gbb(ws, rev_row, cols),
            'profit_margin': gbb(ws, margin_row, cols),
        }

    return launch_manager, brand_manager


def extract_rd(ws, cols):
    """Leadership sheet, rows 13+: one row per named R&D product, format
    'Name - CODE' (a couple are 'CODE - Name' -- handle both)."""
    rd = {}
    r = 13
    while True:
        label = ws.cell(row=r, column=1).value
        if not label or str(label).strip().upper() in ('TOTAL', 'LAUNCH (FIRST 3 MONTHS)'):
            break
        label = str(label).strip()
        parts = [p.strip() for p in label.split('-', 1)]
        if len(parts) == 2:
            a, b = parts
            def looks_like_code(s):
                return bool(re.match(r'^[A-Za-z0-9]+$', s)) and (any(ch.isdigit() for ch in s) or (s.isupper() and len(s) <= 6))
            if looks_like_code(a) and not looks_like_code(b):
                code = a
            elif looks_like_code(b):
                code = b
            else:
                code = b
        else:
            code = label
        rd[code] = {'label': label, 'revenue': gbb(ws, r, cols), 'profit_margin': None}
        r += 1
    return rd


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--bm', required=True, help='Path to the BM Scorecard xlsx')
    ap.add_argument('--leadership', required=True, help='Path to the Leadership Scorecard xlsx')
    ap.add_argument('--out-dir', default='public/targets_monthly')
    ap.add_argument('--year', type=int, default=2026)
    args = ap.parse_args()

    bm_wb = load_workbook(args.bm, data_only=True)
    bm_ws = bm_wb['Brands']
    bm_month_cols = find_month_columns(bm_ws)

    ls_wb = load_workbook(args.leadership, data_only=True)
    ls_ws = ls_wb['Leadership']
    ls_month_cols = find_month_columns(ls_ws)

    written, skipped = [], []
    for month_num in range(1, 13):
        month_str = f'{args.year}-{month_num:02d}'
        if month_num not in bm_month_cols or month_num not in ls_month_cols:
            skipped.append((month_str, 'not present in one of the source sheets'))
            continue

        launch_manager, brand_manager = extract_bm_scorecard(bm_ws, bm_month_cols[month_num])
        rd_team = extract_rd(ls_ws, ls_month_cols[month_num])

        # Gate on real data using Germany's own revenue triple (Launch
        # Manager) -- if the whole month is still a live-tracker
        # placeholder, Germany's revenue is 0/0/0 same as everything else.
        if not has_real_data(launch_manager['germany']['revenue']):
            skipped.append((month_str, 'placeholder/zero data (month not yet reached in the live tracker)'))
            continue

        result = {
            'month': month_str,
            'note': 'Real per-month Good/Better/Best targets, sourced directly from the BM Scorecard (Brands sheet: Launch Mgr, Brand Mgr) and the Leadership Scorecard (Leadership sheet: R&D, revenue only).',
            'rd_team': rd_team,
            'launch_manager': launch_manager,
            'brand_manager': brand_manager,
        }
        out_path = f'{args.out_dir}/{month_str}.json'
        with open(out_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, indent=1, ensure_ascii=False, default=str)
        written.append(month_str)

    print('Written:', written)
    print('Skipped:', skipped)


if __name__ == '__main__':
    main()
